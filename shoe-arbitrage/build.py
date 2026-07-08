#!/usr/bin/env python3
"""
신발 아비트리지 보드 빌드 파이프라인.

  1) 엑셀 넣기(ingest): 포이즌×리테일 엑셀 → data/<조사일>.json 스냅샷으로 정규화
  2) 빌드(build):       data/*.json 스냅샷 전부 → board.html (자체 완결형)

사용법
  python build.py ingest data/raw/2026-07-04.xlsx     # 날짜 자동 감지(요약 시트)
  python build.py ingest some.xlsx --date 2026-07-11  # 날짜 직접 지정
  python build.py build                               # board.html 다시 생성
  python build.py ingest some.xlsx                    # ingest 후 자동으로 build까지

의존성:  pip install openpyxl
"""
import sys, os, json, re, glob, argparse, shutil
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(HERE, "data")
RAW_DIR = os.path.join(DATA_DIR, "raw")
TEMPLATE = os.path.join(HERE, "template.html")
OUTPUT = os.path.join(HERE, "board.html")

# 전체매칭 시트 컬럼 순서
# 0 순번 1 코드 2 상품명 3 브랜드 4 사이트 5 소싱가 6 할인율 7 재고
# 8 포이즌시장가 9 거래량 10 마진 11 마진율 12 발매일 13 조사일 ...


def _pick_sheet(wb, prefix):
    for name in wb.sheetnames:
        if name.startswith(prefix):
            return wb[name]
    return None


def _detect_date(wb, fallback):
    ws = _pick_sheet(wb, "요약")
    if ws:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str):
                    m = re.search(r"(20\d{2})[-.](\d{1,2})[-.](\d{1,2})", cell)
                    if m:
                        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    return fallback


def normalize(xlsx_path, date=None):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = _pick_sheet(wb, "전체매칭") or _pick_sheet(wb, "수요")
    if ws is None:
        raise SystemExit("!! '전체매칭' 시트를 찾을 수 없습니다: " + xlsx_path)

    fallback = re.search(r"(20\d{2})-?(\d{2})-?(\d{2})", os.path.basename(xlsx_path))
    fb_date = f"{fallback.group(1)}-{fallback.group(2)}-{fallback.group(3)}" if fallback else "0000-00-00"
    date = date or _detect_date(wb, fb_date)

    rows = list(ws.iter_rows(values_only=True))
    brands, sites = [], []

    def idx(lst, v):
        v = v or "기타"
        if v not in lst:
            lst.append(v)
        return lst.index(v)

    recs = []
    for r in rows[1:]:
        if r[0] is None:
            continue
        name = (r[2] or "").strip()
        recs.append([
            name, idx(brands, r[3]), idx(sites, r[4]),
            int(r[5] or 0), int(r[6] or 0), int(r[8] or 0),
            int(r[9] or 0), int(r[10] or 0), round(r[11] or 0, 1),
            str(r[1] or "").strip(),   # 9: 상품 코드 (요청 참조용)
        ])

    pos = [x for x in recs if x[7] > 0]
    summary = {
        "total": len(recs),
        "arb": len(pos),
        "avgMargin": round(sum(x[7] for x in pos) / len(pos)) if pos else 0,
        "avgMpct": round(sum(x[8] for x in pos) / len(pos), 1) if pos else 0,
        "bestMargin": max((x[7] for x in recs), default=0),
        "date": date,
    }
    return date, {"b": brands, "s": sites, "r": recs, "m": summary}


def ingest(xlsx_path, date=None):
    date, snap = normalize(xlsx_path, date)
    os.makedirs(DATA_DIR, exist_ok=True)
    os.makedirs(RAW_DIR, exist_ok=True)
    out = os.path.join(DATA_DIR, f"{date}.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(snap, f, ensure_ascii=False, separators=(",", ":"))
    # 원본 보존(같은 파일이면 skip)
    raw_dest = os.path.join(RAW_DIR, f"{date}.xlsx")
    if os.path.abspath(xlsx_path) != os.path.abspath(raw_dest):
        shutil.copy2(xlsx_path, raw_dest)
    print(f"[ingest] {date}: 매칭 {snap['m']['total']}건 · 마진＋ {snap['m']['arb']}건 → {os.path.relpath(out, HERE)}")
    return date


def _peek_date(xlsx):
    """전체 파싱 없이 조사일만 추정 (파일명 → 요약시트 순)."""
    m = re.search(r"(20\d{2})-?(\d{2})-?(\d{2})", os.path.basename(xlsx))
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    try:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
        d = _detect_date(wb, None)
        wb.close()
        return d
    except Exception:
        return None


def sync(from_dir=None, pattern="*.xlsx"):
    """새 조사일 엑셀을 전부 ingest 후 build.
    - data/raw/ 안의 엑셀은 항상 확인
    - from_dir 지정 시 그 폴더(감시 폴더)의 새 엑셀도 끌어와 처리 (매일 크론용)
    """
    done = {os.path.splitext(os.path.basename(p))[0] for p in glob.glob(os.path.join(DATA_DIR, "*.json"))}
    added = 0

    candidates = []
    if from_dir:
        candidates += sorted(glob.glob(os.path.join(os.path.expanduser(from_dir), pattern)))
    candidates += sorted(glob.glob(os.path.join(RAW_DIR, "*.xlsx")))

    for xlsx in candidates:
        date = _peek_date(xlsx)
        if date and date in done:
            continue  # 이미 반영된 조사일 — 건너뜀
        try:
            date = ingest(xlsx, date)          # raw로 복사 + 스냅샷 저장
            done.add(date)
            added += 1
        except Exception as e:
            print(f"[skip]   {os.path.basename(xlsx)} — {e}")

    if added == 0:
        print("[sync]   새로 추가할 엑셀 없음 — 최신 상태.")
    build()


def build():
    files = sorted(glob.glob(os.path.join(DATA_DIR, "*.json")))
    if not files:
        raise SystemExit("!! data/*.json 스냅샷이 없습니다. 먼저 `python build.py ingest <엑셀>` 실행.")
    snaps = {}
    for fp in files:
        date = os.path.splitext(os.path.basename(fp))[0]
        with open(fp, encoding="utf-8") as f:
            snaps[date] = json.load(f)
    dates = sorted(snaps.keys())
    payload = {"dates": dates, "latest": dates[-1], "snaps": snaps,
               "buildAt": datetime.now().strftime("%Y-%m-%d %H:%M")}
    with open(TEMPLATE, encoding="utf-8") as f:
        html = f.read()
    blob = "<script>window.DATA=" + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";</script>"
    html = html.replace("<!--DATA-->", blob)
    with open(OUTPUT, "w", encoding="utf-8") as f:
        f.write(html)
    kb = round(os.path.getsize(OUTPUT) / 1024)
    print(f"[build]  스냅샷 {len(dates)}개 ({', '.join(dates)}) → {os.path.relpath(OUTPUT, HERE)} ({kb} KB)")


def main():
    ap = argparse.ArgumentParser(description="신발 아비트리지 보드 빌드")
    sub = ap.add_subparsers(dest="cmd")
    pi = sub.add_parser("ingest", help="엑셀을 스냅샷 JSON으로 정규화")
    pi.add_argument("xlsx")
    pi.add_argument("--date", help="조사일 YYYY-MM-DD (미지정 시 자동 감지)")
    pi.add_argument("--no-build", action="store_true", help="ingest만 하고 build 생략")
    sub.add_parser("build", help="스냅샷들로 board.html 생성")
    ps = sub.add_parser("sync", help="새 엑셀 자동 ingest + build (크론용)")
    ps.add_argument("--from", dest="from_dir", help="감시 폴더 — 이 폴더의 새 엑셀도 끌어옴")
    ps.add_argument("--pattern", default="*.xlsx", help="감시 폴더 파일 패턴 (기본 *.xlsx)")
    args = ap.parse_args()

    if args.cmd == "ingest":
        ingest(args.xlsx, args.date)
        if not args.no_build:
            build()
    elif args.cmd == "build":
        build()
    elif args.cmd == "sync":
        sync(args.from_dir, args.pattern)
    else:
        ap.print_help()


if __name__ == "__main__":
    main()
